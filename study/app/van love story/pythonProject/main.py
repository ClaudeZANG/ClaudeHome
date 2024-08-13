import tkinter as tk
from tkinter import ttk
from docx import Document
import random

# 替换文档中的占位符函数
def replace_placeholders_with_choices(doc_path, output_path, user_choices):
    doc = Document(doc_path)
    for paragraph in doc.paragraphs:
        for placeholder, choice in user_choices.items():
            if placeholder in paragraph.text:
                paragraph.text = paragraph.text.replace(placeholder, choice)
    # 保存更新后的文档到指定的输出路径
    doc.save(output_path)

# 根据用户选择更新并替换文档
def update_document():
    user_choices = {
        "XX": weather_var.get(),
        "QQ": name_var.get(),
        "W": gender_var.get(),
    }
    replace_placeholders_with_choices(doc_path, output_path, user_choices)
    window.quit()

# 创建Tkinter窗口
window = tk.Tk()
window.title("选择替换选项")

# 创建并设置变量来存储用户选择
weather_var = tk.StringVar()
name_var = tk.StringVar()
gender_var = tk.StringVar()

# 创建下拉菜单让用户选择
weather_label = ttk.Label(window, text="选择天气:")
weather_label.pack()
weather_dropdown = ttk.Combobox(window, textvariable=weather_var)
weather_dropdown['values'] = ("晴朗", "下雨", "多云")
weather_dropdown.pack()

name_label = ttk.Label(window, text="选择姓名:")
name_label.pack()
name_dropdown = ttk.Combobox(window, textvariable=name_var)
name_dropdown['values'] = ("小明", "小红", "小刚")
name_dropdown.pack()

gender_label = ttk.Label(window, text="选择性别:")
gender_label.pack()
gender_dropdown = ttk.Combobox(window, textvariable=gender_var)
gender_dropdown['values'] = ("男", "女")
gender_dropdown.pack()

# 创建按钮，用户点击后将根据选择更新文档
update_button = ttk.Button(window, text="更新文档", command=update_document)
update_button.pack()

# 模板文档路径
doc_path = 'C:\\Users\\chuan\\OneDrive\\s\\app\\van love story\\pythonProject\\dist\\template_document.docx'
# 输出文档的路径和名称
output_path = 'C:\\Users\\chuan\\OneDrive\\s\\app\\van love story\\pythonProject\\dist\\my_updated_document.docx'

window.mainloop()
